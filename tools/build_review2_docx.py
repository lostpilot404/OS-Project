"""Build the editable Review 2 DOCX from the committed artifacts.

    python tools/build_review2_docx.py                     # write the canonical DOCX
    python tools/build_review2_docx.py --output /tmp/x.docx

Same source as the PDF: the narrative comes from :mod:`review2_content` and the figures
from :mod:`review2_figures`, so the two deliverables cannot disagree. The DOCX is the
editable submission copy — it carries the same tables and figures as the PDF, plus the
Mermaid source of every diagram so a reader can re-render or edit them.
"""

from __future__ import annotations

import argparse
import datetime as _datetime
import re
import sys
from pathlib import Path
from typing import Dict, Sequence, Tuple

REPO_ROOT = Path(__file__).resolve().parent.parent
TOOLS_DIR = Path(__file__).resolve().parent
for _path in (str(REPO_ROOT), str(TOOLS_DIR)):
    if _path not in sys.path:
        sys.path.insert(0, _path)

from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.shared import Cm, Pt, RGBColor  # noqa: E402

import review2_content as content  # noqa: E402
import review2_figures as figures  # noqa: E402
from review2_data import Artifacts  # noqa: E402
from review2_layout import REPORT_TITLE, REPO_URL, REVIEW_LINE  # noqa: E402

DEFAULT_OUTPUT = REPO_ROOT / "docs" / "Review_2_Report.docx"
DEFAULT_FIGURES = REPO_ROOT / "figures" / "review2"

MONO = "Consolas"
TAG_RE = re.compile(r"</?(?:b|i|u|font|sub|sup)[^>]*>")
BOLD_RE = re.compile(r"<b>(.*?)</b>", re.DOTALL)
ITALIC_RE = re.compile(r"<i>(.*?)</i>", re.DOTALL)
CODE_RE = re.compile(r"<font face='Courier'>(.*?)</font>", re.DOTALL)


def new_document() -> Document:
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    for name, size in (("Heading 1", 15), ("Heading 2", 12.5), ("Heading 3", 11.5),
                       ("Title", 20)):
        style = document.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
    return document


def _runs(paragraph, text: str) -> None:
    """Write inline markup (bold, italic, monospace) into one paragraph."""
    text = CODE_RE.sub(r"`\1`", text)
    tokens = re.split(r"(<b>.*?</b>|<i>.*?</i>|`[^`]*`)", text, flags=re.DOTALL)
    for token in tokens:
        if not token:
            continue
        bold = italic = mono = False
        body = token
        if token.startswith("<b>"):
            bold, body = True, token[3:-4]
        elif token.startswith("<i>"):
            italic, body = True, token[3:-4]
        elif token.startswith("`"):
            mono, body = True, token[1:-1]
        body = TAG_RE.sub("", body)
        run = paragraph.add_run(body)
        run.bold = bold
        run.italic = italic
        if mono:
            run.font.name = MONO
            run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0, 0, 0)


def para(document: Document, text: str) -> None:
    _runs(document.add_paragraph(), text)


def bullets(document: Document, items: Sequence[str]) -> None:
    for item in items:
        _runs(document.add_paragraph(style="List Bullet"), item)


def numbers(document: Document, items: Sequence[str]) -> None:
    for index, item in enumerate(items, start=1):
        _runs(document.add_paragraph(style="List Number"), item)


def caption(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0, 0, 0)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def code(document: Document, text: str) -> None:
    for line in text.strip("\n").split("\n"):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.left_indent = Pt(14)
        run = paragraph.add_run(line if line else " ")
        run.font.name = MONO
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(0, 0, 0)
    document.add_paragraph().paragraph_format.space_after = Pt(4)


def table(document: Document, rows: Sequence[Sequence], widths: Sequence[float]) -> None:
    table_obj = document.add_table(rows=0, cols=len(rows[0]))
    table_obj.style = "Table Grid"
    available = 16.5  # cm, inside the margins of an A4 page
    for row_index, row in enumerate(rows):
        cells = table_obj.add_row().cells
        for index, (cell, value) in enumerate(zip(cells, row)):
            cell.text = ""
            _runs(cell.paragraphs[0], str(value))
            if row_index == 0:
                for run in cell.paragraphs[0].runs:
                    run.bold = True
            for run in cell.paragraphs[0].runs:
                run.font.size = Pt(8.5)
            if index < len(widths):
                cell.width = Cm(available * float(widths[index]))
    document.add_paragraph().paragraph_format.space_after = Pt(4)


def image(document: Document, path: Path, caption_text: str) -> None:
    document.add_picture(str(path), width=Cm(16.0))
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption(document, caption_text)


# --------------------------------------------------------------------------------------
# rendering
# --------------------------------------------------------------------------------------
def render_block(block: Tuple, document: Document, fig_paths: Dict[str, Path]) -> None:
    kind, payload = block[0], block[1]
    if kind == "h1":
        document.add_heading(payload, level=1)
    elif kind == "h2":
        document.add_heading(payload, level=2)
    elif kind == "p":
        para(document, payload)
    elif kind == "bullets":
        bullets(document, payload)
    elif kind == "numbers":
        numbers(document, payload)
    elif kind == "code":
        code(document, payload["text"])
        if payload.get("caption"):
            caption(document, payload["caption"])
    elif kind == "table":
        table(document, payload["rows"], payload["widths"])
        if payload.get("caption"):
            caption(document, payload["caption"])
    elif kind == "kpi":
        bullets(document, [f"{value.replace('<br/>', ' ')} — {label.replace('<br/>', ' ')}"
                           for value, label in payload])
    elif kind == "fig":
        path = fig_paths.get(payload["key"])
        if path and path.exists():
            image(document, path, payload["caption"])
            if payload.get("mermaid"):
                code(document, "```mermaid\n" + payload["mermaid"].strip() + "\n```")
        else:  # pragma: no cover
            para(document, f"[figure {payload['key']} unavailable]")
    else:  # pragma: no cover
        raise ValueError(f"unknown block kind: {kind!r}")


def cover(document: Document, data: Artifacts, generated: str, test_line: str) -> None:
    fields = content.cover_fields(data, test_line, generated)
    document.add_heading(REPORT_TITLE, level=0)
    for line in (REVIEW_LINE, "Self-healing scheduling in a discrete-event CPU simulator",
                 f"Repository: {REPO_URL}   ·   Report generated {generated}"):
        paragraph = document.add_paragraph()
        run = paragraph.add_run(line)
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0, 0, 0)
    document.add_heading("Cover details", level=2)
    table(document, [["Field", "Value"]] + fields["meta"], [0.24, 0.76])
    document.add_heading("Summary", level=2)
    para(document, fields["abstract"])
    document.add_heading("Key figures", level=2)
    bullets(document, [f"{value.replace('<br/>', ' ')} — {label.replace('<br/>', ' ')}"
                       for value, label in fields["kpis"]])


def contents(document: Document, blocks: Sequence[Tuple]) -> None:
    document.add_heading("Contents", level=1)
    para(document, "The report follows the Review 2 rubric item by item: the nine numbered "
                   "headings are the nine required items, followed by an appendix of "
                   "reproduction commands, seeds and preserved workflows. This list is a "
                   "reading aid; the DOCX also carries a live navigation pane.")
    for level, text in content.outline(blocks):
        paragraph = document.add_paragraph()
        run = paragraph.add_run(text if level == 0 else "    " + text)
        run.bold = level == 0
        run.font.size = Pt(10 if level == 0 else 9.5)
        run.font.color.rgb = RGBColor(0, 0, 0)


def build(root: Path = REPO_ROOT, output: Path = DEFAULT_OUTPUT,
          figures_dir: Path = DEFAULT_FIGURES, with_tests: bool = True,
          with_dataset: bool = True, docs_copy: Path | None = None) -> Path:
    """Render the DOCX to ``output`` and return the path written."""
    from build_review2_report import collect_test_counts  # noqa: PLC0415

    generated = _datetime.date.today().isoformat()
    test_line = (collect_test_counts(root)[0] if with_tests
                 else "test suite not re-run for this build")
    data = Artifacts(root, with_dataset=with_dataset)
    figures_dir.mkdir(parents=True, exist_ok=True)
    fig_paths = figures.render_all(data, content.MERMAID, figures_dir)

    document = new_document()
    cover(document, data, generated, test_line)
    document.add_page_break()
    blocks = content.sections(data)
    contents(document, blocks)
    document.add_page_break()
    for block in blocks:
        render_block(block, document, fig_paths)

    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output))
    if docs_copy:
        Path(docs_copy).parent.mkdir(parents=True, exist_ok=True)
        Path(docs_copy).write_bytes(output.read_bytes())
    return output


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Build the Review 2 report DOCX.")
    parser.add_argument("--root", type=Path, default=REPO_ROOT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--figures-dir", type=Path, default=DEFAULT_FIGURES)
    parser.add_argument("--docs-copy", type=Path, default=None)
    parser.add_argument("--skip-tests", action="store_true")
    parser.add_argument("--no-dataset", action="store_true")
    args = parser.parse_args(argv)
    path = build(root=args.root, output=args.output, figures_dir=args.figures_dir,
                 with_tests=not args.skip_tests, with_dataset=not args.no_dataset,
                 docs_copy=args.docs_copy)
    print(f"wrote {path} ({path.stat().st_size / 1024:.0f} KiB)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
